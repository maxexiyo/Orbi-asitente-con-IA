"""Crear Excel (.xlsx) y PowerPoint (.pptx) DIRECTO, sin abrir ningún programa ni teclear (igual que redactar_documento con Word).

  escribir_xlsx(ruta, contenido, agregar=False)            -> {resultado, ruta, hojas} | {error}
  escribir_pptx(ruta, contenido, tema="azul", agregar=False) -> {resultado, ruta, diapositivas} | {error}

`contenido` es markdown simple (lo redacta otro modelo; ver live_orb.redactar con formato 'hoja' / 'presentacion'):

  EXCEL                                   POWERPOINT
  ## Hoja: Gastos                         # Título de la presentación
  | Concepto | Monto |                   Subtítulo (líneas de texto debajo del título)
  |---|---|
  | Renta | $5,000 |                     ## Título de la diapositiva
  | Total | =SUM(B2:B2) |                - punto
  [grafica: barras]  (antes de la tabla) ![descripción](imagen.png)
                                         [grafica: pastel] + tabla  |  tabla normal
                                         > Notas: lo que dirá el orador

Necesita:  pip install openpyxl python-pptx   (openpyxl ya se usaba para leer Excel).
Igual que el Word: NUNCA pisa un archivo (crea «nombre (2)»); `agregar=True` añade hojas / diapositivas a uno existente que NO esté abierto.
"""
import datetime as _dt
import io, os, re
from pathlib import Path

import archivos_ops as AO
from archivos_ops import OpError

# --------------------------------------------------------------------------- comunes
_TABLA_SEP = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")
_GRAFICA = re.compile(r"^\s*\[\s*gr[aá]fica?s?\s*[:\-]?\s*([a-záéíóúñ ]*)\]\s*(.*)$", re.I)
_IMG = re.compile(r"^\s*!\[([^\]]*)\]\(([^)]+)\)\s*$")
_TIPOS_GRAFICA = {"barra": "barras", "barras": "barras", "columna": "barras", "columnas": "barras", "linea": "lineas", "línea": "lineas",
                  "lineas": "lineas", "líneas": "lineas", "pastel": "pastel", "pay": "pastel", "pie": "pastel", "circular": "pastel",
                  "dona": "pastel"}


def _es_fila_tabla(l: str) -> bool:
    s = l.strip()
    return s.startswith("|") and s.count("|") >= 2


def _celdas(l: str) -> list:
    s = l.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip() for c in s.split("|")]


def _limpia(t: str) -> str:
    """Quita **negritas**, *cursivas* y `código` del markdown, SIN recortar espacios (para los trozos de un párrafo)."""
    t = re.sub(r"\*\*(.+?)\*\*", r"\1", t)
    t = re.sub(r"(?<!\*)\*(?!\s)(.+?)(?<!\s)\*(?!\*)", r"\1", t)
    return re.sub(r"`([^`]*)`", r"\1", t)


def _quitar_md(t: str) -> str:
    return _limpia(t).strip()


def _tipo_grafica(txt: str) -> str:
    for w in re.split(r"\s+", (txt or "").lower().strip()):
        if w in _TIPOS_GRAFICA:
            return _TIPOS_GRAFICA[w]
    return "barras"


def _num(s):
    """'1,200.5' / '$300' / '15%' -> (float, formato). None si no es número."""
    t = str(s).strip()
    m = re.fullmatch(r"(-?)\s*\$\s*(-?)\s*(\d[\d,]*(?:\.\d+)?)", t)
    if m:
        v = float(m.group(3).replace(",", ""))
        return (-v if (m.group(1) or m.group(2)) else v), "$#,##0.00"
    m = re.fullmatch(r"(-?\d+(?:\.\d+)?)\s*%", t)
    if m:
        return float(m.group(1)) / 100, ("0.0%" if "." in m.group(1) else "0%")
    if re.fullmatch(r"-?\d{1,3}(,\d{3})+(\.\d+)?", t):
        return float(t.replace(",", "")), "#,##0.##" if "." in t else "#,##0"
    if re.fullmatch(r"-?\d+(\.\d+)?", t):
        if re.fullmatch(r"-?0\d+", t) or len(t.lstrip("-").split(".")[0]) > 15:     # «007», teléfonos, folios largos: texto
            return None
        return float(t), None
    return None


def _fecha(s):
    t = str(s).strip()
    m = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})", t)
    if m:
        y, mo, d = map(int, m.groups())
    else:
        m = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{4})", t)
        if not m:
            return None
        d, mo, y = map(int, m.groups())
    try:
        return _dt.date(y, mo, d)
    except ValueError:
        return None


def _guardar_seguro(guardar, path: Path) -> None:
    """guardar(ruta_tmp) y luego reemplaza. Si el archivo está abierto en Office, avisa en vez de romperlo."""
    tmp = path.with_name(path.name + ".tmp")
    try:
        guardar(tmp)
        os.replace(tmp, path)
    finally:
        try:
            if tmp.exists():
                tmp.unlink()
        except OSError:
            pass


def _falta(mod: str, pip: str, que: str) -> dict:
    return {"error": f"Para crear {que} falta instalar «{pip}»  (pip install {pip}). Si es el .exe, hay que volver a generarlo con esa librería."}


_ABIERTO = ("Ese archivo está abierto en {app} y no lo puedo modificar por detrás. Ciérralo y lo intento otra vez, "
            "o dime y te lo guardo como uno nuevo.")


# =========================================================================== EXCEL
def _parse_hojas(md: str) -> list:
    """-> [{nombre, bloques:[{tipo:'tabla'|'texto', ...}]}]. `## Hoja: X` o `## X` abre una hoja; sin títulos = una sola."""
    hojas, cur, pend, tabla = [], None, "", []

    def hoja(nombre="Hoja1"):
        nonlocal cur
        if cur is None:
            cur = {"nombre": nombre, "bloques": []}
            hojas.append(cur)
        return cur

    def cerrar_tabla():
        nonlocal tabla, pend
        if tabla:
            filas = [r for r in tabla if not _TABLA_SEP.match("|" + "|".join(r) + "|")]
            if filas:
                hoja()["bloques"].append({"tipo": "tabla", "filas": filas, "grafica": pend})
            pend = ""
        tabla = []

    for raw in (md or "").splitlines():
        l = raw.rstrip()
        m = re.match(r"^\s{0,3}#{1,3}\s+(?:hoja\s*:\s*)?(.+?)\s*#*\s*$", l, re.I)
        if m and not _es_fila_tabla(l):
            cerrar_tabla()
            cur = {"nombre": _quitar_md(m.group(1)), "bloques": []}
            hojas.append(cur)
            continue
        g = _GRAFICA.match(l)
        if g:
            cerrar_tabla()
            pend = _tipo_grafica(g.group(1) or g.group(2))
            continue
        if _es_fila_tabla(l):
            tabla.append(_celdas(l))
            continue
        cerrar_tabla()
        if l.strip():
            hoja()["bloques"].append({"tipo": "texto", "texto": _quitar_md(l.strip().lstrip("-*• ").strip())})
    cerrar_tabla()
    return [h for h in hojas if h["bloques"]]


def _nombre_hoja(wb, nombre: str) -> str:
    n = re.sub(r"[\[\]:*?/\\]", " ", nombre or "Hoja").strip()[:31] or "Hoja"
    base, i = n, 2
    while n in wb.sheetnames:
        suf = f" ({i})"
        n, i = base[:31 - len(suf)] + suf, i + 1
    return n


def _llenar_hoja(ws, bloques: list) -> None:
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as L

    azul, fino = PatternFill("solid", fgColor="2F5BEA"), Side(style="thin", color="C9D3E6")
    borde = Border(left=fino, right=fino, top=fino, bottom=fino)
    fila, anchos, primera = 1, {}, True
    for b in bloques:
        if b["tipo"] == "texto":
            c = ws.cell(row=fila, column=1, value=b["texto"])
            c.font = Font(name="Calibri", bold=True, size=14) if primera else Font(name="Calibri", italic=True, color="555555")
            fila += 1
            continue
        filas = b["filas"]
        ncol = max(len(r) for r in filas)
        ini = fila
        for i, r in enumerate(filas):
            for j in range(ncol):
                txt = r[j] if j < len(r) else ""
                if i == 0:
                    c = ws.cell(row=fila, column=j + 1, value=_quitar_md(txt))
                    c.font, c.fill, c.border = Font(name="Calibri", bold=True, color="FFFFFF"), azul, borde
                    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                else:
                    v, fmt = txt, None
                    if txt.startswith("="):
                        v = txt                                              # fórmula (en inglés: SUM, AVERAGE, IF...)
                    elif (n := _num(txt)) is not None:
                        v, fmt = n
                    elif (f := _fecha(txt)) is not None:
                        v, fmt = f, "dd/mm/yyyy"
                    else:
                        v = _quitar_md(txt)
                    c = ws.cell(row=fila, column=j + 1, value=(v if v != "" else None))
                    if txt.startswith("=") and i > 1 and ws.cell(row=fila - 1, column=j + 1).number_format != "General":
                        fmt = ws.cell(row=fila - 1, column=j + 1).number_format         # un Total hereda el formato (moneda, %) de la columna
                    if fmt:
                        c.number_format = fmt
                    c.border = borde
                    if i % 2 == 0:
                        c.fill = PatternFill("solid", fgColor="F2F6FC")
                    if str(r[0]).strip().lower().startswith(("total", "suma", "promedio")):
                        c.font = Font(name="Calibri", bold=True)
                ancho = len(str(c.value if c.value is not None else txt)) + (4 if i == 0 else 2)
                anchos[j + 1] = max(anchos.get(j + 1, 8), min(ancho if not txt.startswith("=") else 14, 48))
            fila += 1
        fin = fila - 1
        if primera and len(filas) > 1:
            ws.freeze_panes = ws.cell(row=ini + 1, column=1)
        if len(filas) > 2 and not b.get("grafica") and not any(x["tipo"] == "tabla" for x in bloques[:bloques.index(b)]):
            ws.auto_filter.ref = f"A{ini}:{L(ncol)}{fin}"
        if b.get("grafica") and len(filas) > 1 and ncol > 1:
            ch = {"barras": BarChart, "lineas": LineChart, "pastel": PieChart}[b["grafica"]]()
            if b["grafica"] == "barras":
                ch.type = "col"
            nser = 1 if b["grafica"] == "pastel" else ncol - 1
            ch.add_data(Reference(ws, min_col=2, max_col=1 + nser, min_row=ini, max_row=fin), titles_from_data=True)
            ch.set_categories(Reference(ws, min_col=1, min_row=ini + 1, max_row=fin))
            ch.title = _quitar_md(filas[0][1]) if nser == 1 and len(filas[0]) > 1 else None
            ch.height, ch.width = 8, 15
            ws.add_chart(ch, f"{L(ncol + 2)}{ini}")
        fila += 1
        primera = False
    for col, w in anchos.items():
        ws.column_dimensions[L(col)].width = w


def escribir_xlsx(ruta: str, contenido: str, agregar: bool = False) -> dict:
    """Crea un Excel NUEVO (o agrega hojas a uno existente que no esté abierto). `contenido` = markdown con tablas."""
    try:
        import openpyxl
    except ImportError:
        return _falta("openpyxl", "openpyxl", "Excel")
    hojas = _parse_hojas(str(contenido or ""))
    if not hojas:
        return {"error": "No hay datos para la hoja de cálculo (necesito al menos una tabla | a | b |)."}
    try:
        if agregar:
            p = AO.resolver_archivo(ruta)
            if p.suffix.lower() not in (".xlsx", ".xlsm"):
                raise OpError("Solo puedo agregar hojas a archivos .xlsx.")
            wb = openpyxl.load_workbook(p, keep_vba=p.suffix.lower() == ".xlsm")
        else:
            p = AO._destino_nuevo(ruta, ".xlsx")
            if p.suffix.lower() != ".xlsx":
                p = p.with_suffix(".xlsx")
            wb = openpyxl.Workbook()
            wb.remove(wb.active)
        nombres = []
        for h in hojas:
            ws = wb.create_sheet(_nombre_hoja(wb, h["nombre"]))
            _llenar_hoja(ws, h["bloques"])
            nombres.append(ws.title)
        _guardar_seguro(lambda t: wb.save(t), p)
    except OpError as e:
        return {"error": str(e)}
    except PermissionError:
        return {"error": _ABIERTO.format(app="Excel")}
    except Exception as e:  # noqa: BLE001  (archivo dañado, zip raro...)
        return {"error": f"No pude escribir el Excel: {e}"}
    verbo = "Agregué" if agregar else "Guardé"
    res = {"resultado": f"{verbo} «{p.name}» en {p.parent} (hojas: {', '.join(nombres)}).", "ruta": str(p), "hojas": nombres}
    if any(str(c.value).startswith("=") for ws in wb.worksheets if ws.title in nombres for row in ws.iter_rows() for c in row
           if isinstance(c.value, str)):
        res["nota"] = "Tiene fórmulas: Excel las calcula al abrirlo."
    return res


# =========================================================================== POWERPOINT
TEMAS = {   # fondo_oscuro, acento, claro (fondo de diapositivas), texto, texto_suave
    "azul":    ("0B1F3A", "2F7BFF", "F4F7FB", "1B2A41", "5B6B84"),
    "verde":   ("0F2E25", "21A67A", "F3F9F6", "17332B", "5C7A70"),
    "rojo":    ("2B1216", "E5484D", "FBF5F5", "2E1A1D", "7D5A5E"),
    "morado":  ("1E1442", "7C5CFF", "F6F4FD", "261B4D", "6F6590"),
    "naranja": ("2A1A0B", "F08A24", "FDF8F2", "33220F", "86694C"),
    "oscuro":  ("101418", "4DA3FF", "181D23", "ECEFF3", "9AA6B5"),
    "minimal": ("1E293B", "334155", "FFFFFF", "1E293B", "64748B"),
}
_TEMA_ALIAS = {"blue": "azul", "green": "verde", "red": "rojo", "purple": "morado", "morada": "morado", "orange": "naranja",
               "dark": "oscuro", "negro": "oscuro", "oscura": "oscuro", "blanco": "minimal", "simple": "minimal", "sobrio": "minimal"}
MAX_ITEMS = 7          # más viñetas que esto -> sigue en otra diapositiva
MAX_CHARS = 700


def _parse_slides(md: str, portada: bool) -> list:
    """-> [{tipo:'portada'|'diapo', titulo, subtitulo, items:[(tipo,texto,nivel)], tabla, grafica, imagen, notas}]."""
    slides, cur, tabla, pend = [], None, [], ""

    def nueva(titulo, tipo="diapo"):
        nonlocal cur
        cur = {"tipo": tipo, "titulo": _quitar_md(titulo), "subtitulo": [], "items": [], "tabla": None, "grafica": "",
               "imagen": None, "notas": []}
        slides.append(cur)

    def cerrar_tabla():
        nonlocal tabla, pend
        if tabla and cur is not None:
            filas = [r for r in tabla if not _TABLA_SEP.match("|" + "|".join(r) + "|")]
            if filas:
                cur["tabla"], cur["grafica"] = filas, pend
        tabla, pend = [], ""

    for raw in (md or "").splitlines():
        l = raw.rstrip()
        m = re.match(r"^\s{0,3}(#{1,2})\s+(.+?)\s*#*\s*$", l)
        if m and not _es_fila_tabla(l):
            cerrar_tabla()
            nueva(m.group(2), "portada" if (portada and not slides and len(m.group(1)) == 1) else "diapo")
            continue
        if cur is None:
            if not l.strip():
                continue
            nueva("Presentación", "portada" if portada else "diapo")
        g = _GRAFICA.match(l)
        if g:
            cerrar_tabla()
            pend = _tipo_grafica(g.group(1) or g.group(2))
            continue
        if _es_fila_tabla(l):
            tabla.append(_celdas(l))
            continue
        cerrar_tabla()
        if not l.strip():
            continue
        mi = _IMG.match(l)
        if mi:
            cur["imagen"] = (mi.group(2).strip(), mi.group(1).strip())
            continue
        mn = re.match(r"^\s*(?:>\s*)?(?:notas?|nota del orador|speaker notes?)\s*:\s*(.*)$", l, re.I)
        if mn or l.lstrip().startswith(">"):
            cur["notas"].append((mn.group(1) if mn else l.lstrip().lstrip(">").strip()))
            continue
        if cur["tipo"] == "portada":
            cur["subtitulo"].append(_quitar_md(l.strip().lstrip("-*• ").strip()))
            continue
        mh = re.match(r"^\s*#{3,6}\s+(.+)$", l)
        if mh:
            cur["items"].append(("sub", _quitar_md(mh.group(1)), 0))
            continue
        mb = re.match(r"^(\s*)(?:[-*•]|\d+[.)])\s+(.+)$", l)
        if mb:
            cur["items"].append(("punto", mb.group(2).strip(), 1 if len(mb.group(1).replace("\t", "  ")) >= 2 else 0))
        else:
            cur["items"].append(("texto", l.strip(), 0))
    cerrar_tabla()
    # diapositivas demasiado llenas -> se parten (sin tabla ni imagen, que ya ocupan sitio)
    out = []
    for s in slides:
        its = s["items"]
        if s["tipo"] == "diapo" and len(its) > MAX_ITEMS and not s["tabla"]:
            partes = -(-len(its) // MAX_ITEMS)
            por = -(-len(its) // partes)                        # parejas: 9 viñetas -> 5 + 4, no 7 + 2
            trozos = [its[i:i + por] for i in range(0, len(its), por)]
            for k, t in enumerate(trozos):
                out.append({**s, "items": t, "titulo": s["titulo"] + (" (cont.)" if k else ""), "imagen": s["imagen"] if k == 0 else None,
                            "notas": s["notas"] if k == 0 else []})
        else:
            out.append(s)
    return out


def _rgb(h):
    from pptx.dml.color import RGBColor
    return RGBColor.from_string(h)


def _mezcla(a: str, b: str, t: float) -> str:
    """Color entre a y b (t=0 -> a, t=1 -> b)."""
    ca, cb = [int(a[i:i + 2], 16) for i in (0, 2, 4)], [int(b[i:i + 2], 16) for i in (0, 2, 4)]
    return "".join(f"{round(x + (y - x) * t):02X}" for x, y in zip(ca, cb))


def _lum(h: str) -> float:
    r, g, b = [int(h[i:i + 2], 16) / 255 for i in (0, 2, 4)]
    return 0.299 * r + 0.587 * g + 0.114 * b


def _paleta_graf(acento: str, n: int) -> list:
    """Colores para gráficas: el acento del tema y sus vecinos de tono (se ven armónicos juntos)."""
    import colorsys
    r, g, b = [int(acento[i:i + 2], 16) / 255 for i in (0, 2, 4)]
    h, l, sat = colorsys.rgb_to_hls(r, g, b)
    sat = max(sat, 0.35)
    out = [acento]
    for dh in (0.08, -0.08, 0.16, -0.16, 0.25, -0.25):
        rr, gg, bb = colorsys.hls_to_rgb((h + dh) % 1, min(max(l, 0.38), 0.58), sat)
        out.append(f"{int(rr * 255):02X}{int(gg * 255):02X}{int(bb * 255):02X}")
    return (out * 3)[:max(n, 1)]


def _paleta(t: tuple) -> dict:
    oscuro, acento, claro, txt, suave = t
    es_claro = _lum(claro) > 0.5
    return dict(oscuro=oscuro, acento=acento, claro=claro, texto=txt, suave=suave,
                tarjeta="FFFFFF" if es_claro else _mezcla(claro, "FFFFFF", 0.07),
                borde=_mezcla(claro, txt, 0.13), halo=_mezcla(acento, claro, 0.88), es_claro=es_claro)


def _runs(par, texto, size, color, bold=False, font="Calibri", italic=False):
    from pptx.util import Pt
    for i, trozo in enumerate(re.split(r"\*\*(.+?)\*\*", texto)):
        if not trozo:
            continue
        r = par.add_run()
        r.text = _limpia(trozo) if i % 2 == 0 else trozo
        r.font.size, r.font.bold, r.font.name = Pt(size), bool(bold or i % 2 == 1), font
        r.font.italic = italic
        r.font.color.rgb = _rgb(color)


def _forma(slide, tipo, x, y, w, h, color, borde=None, radio=None):
    from pptx.enum.shapes import MSO_SHAPE
    s = slide.shapes.add_shape(tipo, int(x), int(y), int(w), int(h))
    s.fill.solid()
    s.fill.fore_color.rgb = _rgb(color)
    if borde:
        s.line.color.rgb = _rgb(borde)
        s.line.width = 9525
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    st = s._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}style")
    if st is not None:                                  # sin sombra del tema (LibreOffice y otros visores la ponen aunque el efecto esté vacío)
        er = st.find("{http://schemas.openxmlformats.org/drawingml/2006/main}effectRef")
        if er is not None:
            er.set("idx", "0")
    if radio is not None and tipo == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radio
    return s


def _rect(slide, x, y, w, h, color):
    from pptx.enum.shapes import MSO_SHAPE
    return _forma(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, color)


def _texto(slide, x, y, w, h, anchor="top", margen=True):
    from pptx.enum.text import MSO_ANCHOR
    from pptx.util import Pt
    tb = slide.shapes.add_textbox(int(x), int(y), int(w), int(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[anchor]
    if not margen:
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Pt(0)
    return tb, tf


def _tam_ajustado(parrafos: list, ancho_emu: float, alto_emu: float, maximo: float, minimo: float = 14) -> float:
    """Tamaño de letra más grande (entre maximo y minimo) con el que el texto cabe en la caja (estimado por caracteres)."""
    ancho_pt, alto_pt = ancho_emu / 12700, alto_emu / 12700
    tam = maximo
    while tam > minimo:
        lineas = sum(max(1, -(-int(len(p) * 0.53 * tam) // max(int(ancho_pt), 1))) for p in parrafos)
        if lineas * tam * 1.22 + max(len(parrafos) - 1, 0) * tam * 0.35 <= alto_pt:
            break
        tam -= 1
    return tam


def _imagen_bytes(ref):
    """ref: ruta (Path) ya validada, o texto (archivo del usuario: se valida con las carpetas permitidas)."""
    from PIL import Image
    p = ref if isinstance(ref, Path) else AO.resolver_archivo(str(ref).strip().strip('"'))
    if p.is_dir():
        raise OpError(f"«{p.name}» es una carpeta, no una imagen.")
    try:
        im = Image.open(p)
        im.load()
    except Exception:  # noqa: BLE001
        raise OpError(f"No pude abrir la imagen «{p.name}» (¿es png, jpg, gif, bmp o webp?).")
    if p.suffix.lower() in (".png", ".jpg", ".jpeg", ".gif", ".bmp"):
        return io.BytesIO(p.read_bytes()), im.size
    buf = io.BytesIO()
    (im if im.mode in ("RGB", "RGBA", "L", "LA", "P") else im.convert("RGBA")).save(buf, "PNG")
    buf.seek(0)
    return buf, im.size


def _poner_imagen(slide, ref, x, y, w, h, redondo=True, cubrir=True):
    """Pone la imagen llenando la caja (recortando lo que sobre) y, si redondo, con esquinas redondeadas."""
    from pptx.enum.shapes import MSO_SHAPE
    from lxml import etree
    data, (pw, ph) = _imagen_bytes(ref)
    if cubrir:
        pic = slide.shapes.add_picture(data, int(x), int(y), int(w), int(h))
        ar_img, ar_box = pw / ph, w / h
        if ar_img > ar_box:
            c = (1 - ar_box / ar_img) / 2
            pic.crop_left = pic.crop_right = c
        else:
            c = (1 - ar_img / ar_box) / 2
            pic.crop_top = pic.crop_bottom = c
    else:
        k = min(w / pw, h / ph)
        iw, ih = int(pw * k), int(ph * k)
        pic = slide.shapes.add_picture(data, int(x + (w - iw) / 2), int(y + (h - ih) / 2), iw, ih)
    if redondo:
        pic.auto_shape_type = MSO_SHAPE.ROUNDED_RECTANGLE
        pg = pic._element.spPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}prstGeom")
        if pg is not None:
            av = pg.find("{http://schemas.openxmlformats.org/drawingml/2006/main}avLst")
            if av is None:
                av = etree.SubElement(pg, "{http://schemas.openxmlformats.org/drawingml/2006/main}avLst")
            for g_ in list(av):
                av.remove(g_)
            gd = etree.SubElement(av, "{http://schemas.openxmlformats.org/drawingml/2006/main}gd")
            gd.set("name", "adj")
            gd.set("fmla", "val 5500")
    return pic


def _pie_imagen(sl, im, x, y, w, P):
    """Texto chico bajo la imagen: descripción y crédito (las fotos de internet traen autor y licencia)."""
    alt, cred = (im[1] if len(im) > 1 else ""), (im[2] if len(im) > 2 else "")
    txt = " — ".join(t for t in (alt, cred) if t)[:150]
    if not txt:
        return
    _, tf = _texto(sl, x, y, w, 330000, "top", margen=False)
    _runs(tf.paragraphs[0], txt, 10.5, P["suave"], italic=True)


def _poner_tabla(slide, filas, x, y, w, h, P):
    from pptx.util import Pt
    ncol, nf = max(len(r) for r in filas), len(filas)
    alto_f = min(int(h / nf), 720000)
    gf = slide.shapes.add_table(nf, ncol, int(x), int(y), int(w), alto_f * nf)
    t = gf.table
    tam = 22 if nf <= 6 and ncol <= 3 else (17 if nf <= 10 else 12)
    for j in range(ncol):
        t.columns[j].width = int(w / ncol)
    for i, r in enumerate(filas):
        t.rows[i].height = alto_f
        for j in range(ncol):
            c = t.cell(i, j)
            c.fill.solid()
            c.fill.fore_color.rgb = _rgb(P["acento"] if i == 0 else (P["tarjeta"] if i % 2 else P["halo"]))
            tf = c.text_frame
            tf.word_wrap = True
            _runs(tf.paragraphs[0], r[j] if j < len(r) else "", tam, "FFFFFF" if i == 0 else P["texto"], bold=(i == 0))
            c.margin_left = c.margin_right = Pt(10)
            c.vertical_anchor = 3                    # centrado


def _poner_grafica(slide, filas, tipo, x, y, w, h, P) -> bool:
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
    from pptx.util import Pt
    if len(filas) < 2 or max(len(r) for r in filas) < 2:
        return False
    ncol = max(len(r) for r in filas)
    series = []
    for j in range(1, 2 if tipo == "pastel" else ncol):
        vals = []
        for r in filas[1:]:
            n = _num(r[j]) if j < len(r) else None
            if n is None:
                return False
            vals.append(n[0])
        series.append((_quitar_md(filas[0][j]) if j < len(filas[0]) else f"Serie {j}", vals))
    cd = CategoryChartData()
    cd.categories = [_quitar_md(r[0]) for r in filas[1:]]
    for nom, vals in series:
        cd.add_series(nom, vals)
    tp = {"barras": XL_CHART_TYPE.COLUMN_CLUSTERED, "lineas": XL_CHART_TYPE.LINE_MARKERS, "pastel": XL_CHART_TYPE.PIE}[tipo]
    ch = slide.shapes.add_chart(tp, int(x), int(y), int(w), int(h), cd).chart
    ch.font.size, ch.font.name = Pt(15), "Calibri"
    ch.font.color.rgb = _rgb(P["texto"])
    ch.has_title = False
    cols = _paleta_graf(P["acento"], max(len(cd.categories) if tipo == "pastel" else len(series), 1))
    ch.has_legend = tipo == "pastel" or len(series) > 1
    if ch.has_legend:
        ch.legend.position, ch.legend.include_in_layout = XL_LEGEND_POSITION.BOTTOM, False
        ch.legend.font.size = Pt(14)
    pl = ch.plots[0]
    if tipo == "pastel":
        for i, pt in enumerate(pl.series[0].points):
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = _rgb(cols[i % len(cols)])
        pl.has_data_labels = True
        pl.data_labels.number_format, pl.data_labels.number_format_is_linked = "General", False
        pl.data_labels.font.size, pl.data_labels.font.bold = Pt(15), True
        pl.data_labels.font.color.rgb = _rgb("FFFFFF")
    else:
        for i, se in enumerate(pl.series):
            if tipo == "barras":
                se.format.fill.solid()
                se.format.fill.fore_color.rgb = _rgb(cols[i % len(cols)])
            else:
                se.format.line.color.rgb = _rgb(cols[i % len(cols)])
                se.format.line.width = Pt(3.5)
                se.smooth = False
        if tipo == "barras":
            pl.gap_width = 55
            pl.has_data_labels = True
            pl.data_labels.number_format, pl.data_labels.number_format_is_linked = "General", False
            pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
            pl.data_labels.font.size = Pt(14)
        va, ca = ch.value_axis, ch.category_axis
        va.major_gridlines.format.line.color.rgb = _rgb(P["borde"])
        va.format.line.fill.background()
        va.tick_labels.font.size = Pt(13)
        ca.tick_labels.font.size = Pt(14)
        ca.format.line.color.rgb = _rgb(P["borde"])
    return True


def _tarjetas(sl, tarjetas: list, x, y, w, h, P) -> None:
    """Una tarjeta por punto: círculo con número a la izquierda y el texto (con sus sub-puntos) a la derecha. Centradas en vertical."""
    from pptx.enum.shapes import MSO_SHAPE
    n = len(tarjetas)
    gap = int(h * 0.028)
    ch = min(int((h - gap * (n - 1)) / n), 1180000)
    y0 = y + (h - (ch * n + gap * (n - 1))) / 2
    d = int(min(ch * 0.58, 560000))
    for i, (texto, subs) in enumerate(tarjetas):
        yy = y0 + i * (ch + gap)
        _forma(sl, MSO_SHAPE.ROUNDED_RECTANGLE, x, yy, w, ch, P["tarjeta"], borde=P["borde"], radio=0.14)
        cx = x + int(ch * 0.2)
        circ = _forma(sl, MSO_SHAPE.OVAL, cx, yy + (ch - d) / 2, d, d, P["acento"])
        tf = circ.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.paragraphs[0].alignment = 2
        _runs(tf.paragraphs[0], str(i + 1), max(14, int(d / 12700 * 0.5)), "FFFFFF", bold=True)
        tx = cx + d + int(ch * 0.22)
        tw = x + w - tx - int(ch * 0.2)
        alto_txt = ch - int(ch * 0.16)
        todo = [texto] + subs
        tam = _tam_ajustado(todo, tw, alto_txt, 28 if not subs else 24, 14)
        _, tf = _texto(sl, tx, yy + int(ch * 0.08), tw, alto_txt, "middle")
        _runs(tf.paragraphs[0], texto, tam, P["texto"])
        for st in subs:
            p = tf.add_paragraph()
            p.space_before = 0
            _runs(p, "– " + st, max(tam - 4, 13), P["suave"])


def _lista(sl, items: list, x, y, w, h, P) -> None:
    """Texto corrido (cuando hay muchas viñetas o hay subtítulos/párrafos): letra grande y centrado en la zona."""
    todo = [i[1] for i in items]
    tam = _tam_ajustado(todo, w - 300000, h, 30, 15)
    _, tf = _texto(sl, x, y, w, h, "middle" if len(items) <= 6 else "top")
    primero = True
    for tipo, t, nivel in items:
        p = tf.paragraphs[0] if primero else tf.add_paragraph()
        primero = False
        p.space_after = Pt_(tam * 0.5)
        if tipo == "sub":
            _runs(p, t, tam + 2, P["acento"], bold=True)
        elif tipo == "punto":
            _runs(p, ("–  " if nivel else "•  ") + t, tam - (4 if nivel else 0), P["suave"] if nivel else P["texto"])
        else:
            _runs(p, t, tam, P["texto"])


def Pt_(v):
    from pptx.util import Pt
    return Pt(v)


def _dibujar(prs, s: dict, pal: tuple, num: int) -> None:
    from pptx.enum.shapes import MSO_SHAPE
    P = _paleta(pal)
    W, H = prs.slide_width, prs.slide_height
    lay = min(prs.slide_layouts, key=lambda l: len(l.placeholders))        # el más vacío (en la plantilla por defecto, «Blank»)
    sl = prs.slides.add_slide(lay)
    for ph in list(sl.placeholders):
        ph._element.getparent().remove(ph._element)
    mx = int(W * 0.065)
    im = s.get("imagen") or None

    def fondo(c):
        sl.background.fill.solid()
        sl.background.fill.fore_color.rgb = _rgb(c)

    def tam_titulo(t, grande, chico):
        return grande if len(t) <= 34 else (int((grande + chico) / 2) if len(t) <= 60 else chico)

    def intentar_imagen(*a, **k):
        try:
            return _poner_imagen(sl, *a, **k)
        except OpError as e:
            s.setdefault("_avisos", []).append(str(e))
            return None

    if s["tipo"] == "portada":
        fondo(P["oscuro"])
        ancho_txt = 0.5 if im else 0.84
        if im:
            if intentar_imagen(im[0], int(W * 0.56), 0, int(W * 0.44), H, redondo=False) is not None:
                _rect(sl, int(W * 0.56), 0, int(W * 0.008), H, P["acento"])
            else:
                im = None
                ancho_txt = 0.84
        if not im:
            _forma(sl, MSO_SHAPE.OVAL, int(W * 0.70), int(-H * 0.18), int(H * 1.15), int(H * 1.15), _mezcla(P["acento"], P["oscuro"], 0.55))
            _forma(sl, MSO_SHAPE.OVAL, int(W * 0.80), int(H * 0.55), int(H * 0.75), int(H * 0.75), P["acento"])
        _rect(sl, 0, 0, int(W * 0.022), H, P["acento"])
        _, tf = _texto(sl, mx, int(H * 0.2), int(W * ancho_txt), int(H * 0.36), "bottom")
        _runs(tf.paragraphs[0], s["titulo"], tam_titulo(s["titulo"], 52, 32), "FFFFFF", bold=True)
        _rect(sl, mx, int(H * 0.6), int(W * 0.12), int(H * 0.009), P["acento"])
        if s["subtitulo"]:
            _, tf = _texto(sl, mx, int(H * 0.64), int(W * ancho_txt), int(H * 0.26))
            for i, ln in enumerate(s["subtitulo"][:4]):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.space_after = Pt_(6)
                _runs(p, ln, 24 if i == 0 else 18, "D8E1F0" if i == 0 else "A3B3CE")
    elif not (s["items"] or s["tabla"] or im):                       # solo título = diapositiva de sección
        fondo(P["acento"])
        _forma(sl, MSO_SHAPE.OVAL, int(W * 0.66), int(-H * 0.25), int(H * 1.3), int(H * 1.3), _mezcla(P["acento"], "FFFFFF", 0.12))
        _forma(sl, MSO_SHAPE.OVAL, int(W * 0.84), int(H * 0.62), int(H * 0.6), int(H * 0.6), _mezcla(P["acento"], "000000", 0.12))
        _, tf = _texto(sl, mx, int(H * 0.3), int(W * 0.72), int(H * 0.4), "middle")
        _runs(tf.paragraphs[0], s["titulo"], tam_titulo(s["titulo"], 50, 32), "FFFFFF", bold=True)
    else:
        fondo(P["claro"])
        _forma(sl, MSO_SHAPE.OVAL, int(W * 0.80), int(-H * 0.30), int(H * 0.62), int(H * 0.62), P["halo"])      # detalle suave arriba a la derecha
        _forma(sl, MSO_SHAPE.ROUNDED_RECTANGLE, mx, int(H * 0.075), int(W * 0.0075), int(H * 0.095), P["acento"], radio=0.5)
        _, tf = _texto(sl, mx + int(W * 0.02), int(H * 0.05), int(W * 0.78), int(H * 0.14), "middle")
        _runs(tf.paragraphs[0], s["titulo"], tam_titulo(s["titulo"], 36, 26), P["texto"], bold=True)
        top, alto, ancho = int(H * 0.235), int(H * 0.685), int(W * 0.87)
        items, tabla = s["items"], s["tabla"]
        puntos = [i for i in items if i[0] == "punto"]
        solo_puntos = bool(items) and len(puntos) == len(items) and items[0][2] == 0
        tarj = []
        if solo_puntos:
            for tipo, t, nivel in items:
                if nivel == 0 or not tarj:
                    tarj.append((t, []))
                else:
                    tarj[-1][1].append(t)
        usa_tarjetas = solo_puntos and 2 <= len(tarj) <= 6 and sum(len(t) + sum(len(x) for x in sb) for t, sb in tarj) < 700

        def poner_texto(x, y, w, h):
            if usa_tarjetas:
                _tarjetas(sl, tarj, x, y, w, h, P)
            elif items:
                _lista(sl, items, x, y, w, h, P)

        if tabla:
            if items and max(len(r) for r in tabla) <= 3:               # texto a la izquierda, tabla/gráfica a la derecha
                wi = int(ancho * 0.40)
                poner_texto(mx, top, wi, alto)
                gx, gw = mx + wi + int(ancho * 0.035), ancho - wi - int(ancho * 0.035)
            elif items:                                                  # tabla ancha: texto arriba (corto) y la tabla debajo
                _lista(sl, items[:3], mx, top, ancho, int(alto * 0.26), P)
                gx, gw, top, alto = mx, ancho, top + int(alto * 0.3), int(alto * 0.7)
            else:
                gx, gw = mx, ancho
            if not (s["grafica"] and _poner_grafica(sl, tabla, s["grafica"], gx, top, gw, alto, P)):
                _poner_tabla(sl, tabla, gx, top, gw, alto, P)
        elif im and items:
            wi = int(ancho * 0.57)
            poner_texto(mx, top, wi, alto)
            ix, iw = mx + wi + int(ancho * 0.04), ancho - wi - int(ancho * 0.04)
            if intentar_imagen(im[0], ix, top, iw, alto - 330000) is not None:
                _pie_imagen(sl, im, ix, top + alto - 300000, iw, P)
        elif im:
            if intentar_imagen(im[0], mx + int(ancho * 0.06), top, int(ancho * 0.88), alto - 330000) is not None:
                _pie_imagen(sl, im, mx + int(ancho * 0.06), top + alto - 300000, int(ancho * 0.88), P)
        else:
            poner_texto(mx, top, ancho, alto)
        _, tf = _texto(sl, int(W * 0.9), int(H * 0.93), int(W * 0.07), int(H * 0.05), "middle")
        tf.paragraphs[0].alignment = 3
        _runs(tf.paragraphs[0], str(num), 12, P["suave"])
    if s["notas"]:
        sl.notes_slide.notes_text_frame.text = "\n".join(s["notas"])


def _tema(nombre: str) -> tuple:
    n = (nombre or "azul").strip().lower()
    return TEMAS.get(_TEMA_ALIAS.get(n, n), TEMAS["azul"])


MAX_IMAGENES = 8


def _clasificar_ref(ref: str):
    """-> ('web'|'ia'|'url'|None, texto). None = archivo local del usuario."""
    m = re.match(r"^\s*(web|ia)\s*:\s*(.+)$", ref or "", re.I)
    if m:
        return m.group(1).lower(), m.group(2).strip()
    if re.match(r"^https?://", ref or "", re.I):
        return "url", ref.strip()
    return None, ref


def _resolver_imagenes(slides: list, resolver, tope_s: float = 85.0) -> dict:
    """Consigue (en paralelo) las imágenes que pidieron como web:/ia:/enlace. Lo que falle se quita sin tirar la presentación."""
    import concurrent.futures as cf
    pend, avisos, puestas = [], [], 0
    for s in slides:
        if not s.get("imagen"):
            continue
        tipo, txt = _clasificar_ref(s["imagen"][0])
        if tipo is None:
            puestas += 1
            continue
        if resolver is None or len(pend) >= MAX_IMAGENES:
            avisos.append(f"«{s['titulo']}»: no pude poner la imagen ({'sin permiso' if resolver is None else 'demasiadas imágenes'}).")
            s["imagen"] = None
            continue
        pend.append((s, tipo, txt))
    if pend:
        ex = cf.ThreadPoolExecutor(max_workers=4)
        futs = {}
        for s, tipo, txt in pend:
            asp = "3:4" if s["tipo"] == "portada" else ("1:1" if s["items"] else "16:9")
            futs[ex.submit(resolver, tipo, txt, s["imagen"][1], asp)] = s
        hechos, _ = cf.wait(futs, timeout=tope_s)
        for f, s in futs.items():
            alt = s["imagen"][1]
            try:
                r = f.result(0) if f in hechos else None
                if not r or not r.get("ruta"):
                    raise RuntimeError((r or {}).get("error") or "tardó demasiado")
                s["imagen"] = (Path(r["ruta"]), alt, r.get("credito", ""))
                puestas += 1
            except Exception as e:                      # noqa: BLE001
                avisos.append(f"«{s['titulo']}»: no conseguí la imagen ({str(e)[:90]}).")
                s["imagen"] = None
        ex.shutdown(wait=False, cancel_futures=True)
    return {"puestas": puestas, "avisos": avisos}


def escribir_pptx(ruta: str, contenido: str, tema: str = "azul", agregar: bool = False, resolver=None) -> dict:
    """Crea una presentación NUEVA (o agrega diapositivas a una existente que no esté abierta). `contenido` = markdown de diapositivas.
    resolver(tipo, texto, alt, aspecto) -> {ruta, credito} consigue las imágenes pedidas como ![](web:…) / ![](ia:…) / ![](https://…)."""
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError:
        return _falta("pptx", "python-pptx", "PowerPoint")
    slides = _parse_slides(str(contenido or ""), portada=not agregar)
    if not slides:
        return {"error": "No hay contenido para la presentación."}
    pal = _tema(tema)
    info_img = _resolver_imagenes(slides, resolver)
    try:
        if agregar:
            p = AO.resolver_archivo(ruta)
            if p.suffix.lower() != ".pptx":
                raise OpError("Solo puedo agregar diapositivas a archivos .pptx.")
            prs = Presentation(str(p))
        else:
            p = AO._destino_nuevo(ruta, ".pptx")
            if p.suffix.lower() != ".pptx":
                p = p.with_suffix(".pptx")
            prs = Presentation()
            prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        n0 = len(prs.slides)
        for i, s in enumerate(slides):
            _dibujar(prs, s, pal, n0 + i + 1)
            info_img["avisos"] += [f"«{s['titulo']}»: {a}" for a in s.get("_avisos", [])]
        _guardar_seguro(lambda t: prs.save(str(t)), p)
    except OpError as e:
        return {"error": str(e)}
    except PermissionError:
        return {"error": _ABIERTO.format(app="PowerPoint")}
    except Exception as e:  # noqa: BLE001
        return {"error": f"No pude escribir la presentación: {e}"}
    verbo = f"Agregué {len(slides)} diapositiva(s) a" if agregar else f"Guardé ({len(slides)} diapositivas)"
    res = {"resultado": f"{verbo} «{p.name}» en {p.parent}.", "ruta": str(p), "diapositivas": len(slides),
           "titulos": [s["titulo"] for s in slides][:20]}
    n_img = sum(1 for s in slides if s.get("imagen") and not s.get("_avisos"))
    if n_img or info_img["avisos"]:
        res["imagenes"] = {"puestas": n_img, "avisos": info_img["avisos"][:8]}
    return res
